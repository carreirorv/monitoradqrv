#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
HUB ADQ - Atalhos do menu e publicação no Netlify

Deixe este script na mesma pasta do index.html do HUB.
Funciona com o HUB (index.html com um menu só), com o Monitor (submenus
Comercial, Estratégico e Auditoria) ou com os dois na mesma pasta (index.html +
monitor.html). O script descobre os menus pelas marcações de cada página e, ao
incluir, excluir, ordenar ou mover, pergunta em qual menu mexer.

Uso manual:      python incluir_relatorios.py
  1. Incluir relatórios no menu
  2. Excluir atalhos do menu
  3. Atualizar relatórios e publicar no Netlify
  4. Atualizar usuários do login a partir da planilha (.xlsx)
  5. Ver os relatórios que estão no menu
  6. Validar se os relatórios estão funcionando
  7. Ordenar os relatórios do menu
  8. Mover relatórios para outro menu (ex.: de Comercial para Auditoria)
  9. Trocar o site do Netlify (quando o endereço do site mudar)
 10. Entrar em outra conta do Netlify (o login abre no Google Chrome)
  0. Sair
  Depois de cada opção o script volta ao menu.

Uso automático:  python incluir_relatorios.py --publicar
  Faz a opção 3 sem perguntas (para o Agendador de Tarefas do Windows).
  Antes de publicar, recarrega os usuários da planilha informada na opção 4.
  O resultado de cada execução fica registrado em publicacao.log.
  Para testar a publicação vendo as mensagens na tela, use testar_publicacao.bat.

Requisitos da publicação: Node.js e Netlify CLI instalados
(npm install -g netlify-cli) e login feito uma vez (netlify login).
Requisito da opção 4: openpyxl (pip install openpyxl).
"""

import datetime
import filecmp
import hashlib
import html
import json
import os
import re
import shutil
import subprocess
import time
import sys
import tempfile
import traceback
import urllib.error
import urllib.request
import unicodedata
from pathlib import Path

INICIO = "<!-- ATALHOS:INICIO -->"
FIM = "<!-- ATALHOS:FIM -->"
PASTA_DESTINO = "relatorios"
PASTA_PUBLICACAO = "_publicar"  # só para limpar a pasta das versões antigas
CONFIG = "hub_config.json"
LOG = "publicacao.log"

ICONE = ('<svg viewBox="0 0 24 24"><path d="M6 3h9l4 4v14H6z"/>'
         '<path d="M15 3v4h4"/><path d="M9 12h6M9 16h6"/></svg>')
CORES = ["", "c2", "c3", "c4", "c5", "c6"]

# Páginas que vão para o site (as que existirem na pasta)
PAGINAS = ["index.html", "monitor.html"]

# Os menus são descobertos pelas marcações de cada página:
#   <!-- ATALHOS:INICIO -->            menu único (HUB ADQ)
#   <!-- ATALHOS:COMERCIAL:INICIO -->  submenu (Monitor), um par para cada submenu
# Assim o script funciona com o HUB, com o Monitor, ou com o Monitor salvo como index.html.
MARCA = re.compile(r"<!-- ATALHOS:(?:([A-Za-z0-9_]+):)?INICIO -->")
ROTULOS = {"COMERCIAL": "Comercial", "ESTRATEGICO": "Estratégico", "AUDITORIA": "Auditoria"}

AUTOMATICO = "--publicar" in sys.argv
PASTA_HUB = Path(__file__).resolve().parent


# ---------------------------------------------------------------- utilidades

def registrar(msg: str) -> None:
    print(msg, flush=True)
    try:
        agora = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with open(PASTA_HUB / LOG, "a", encoding="utf-8") as f:
            f.write(f"[{agora}] {msg}\n")
    except OSError:
        pass


def limpar_caminho(texto: str) -> Path:
    # Remove aspas que o Windows coloca ao arrastar uma pasta para o terminal
    return Path(texto.strip().strip('"').strip("'")).expanduser()


def perguntar(msg: str, padrao: str = "") -> str:
    sufixo = f" [{padrao}]" if padrao else ""
    resposta = input(f"{msg}{sufixo}: ").strip()
    return resposta or padrao


def sim(msg: str, padrao: str = "s") -> bool:
    return perguntar(f"{msg} (s/n)", padrao).lower() in ("s", "sim", "y")


def slug(texto: str) -> str:
    t = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    t = re.sub(r"[^a-zA-Z0-9]+", "-", t).strip("-").lower()
    return t or "relatorio"


def titulo_do_html(arquivo: Path) -> str:
    try:
        trecho = arquivo.read_text(encoding="utf-8", errors="ignore")[:20000]
        m = re.search(r"<title[^>]*>(.*?)</title>", trecho, re.S | re.I)
        if m and m.group(1).strip():
            return html.unescape(m.group(1).strip())
    except OSError:
        pass
    return arquivo.stem.replace("_", " ").replace("-", " ")


def montar_atalho(href: str, titulo: str, descricao: str, cor: str) -> str:
    classe = f"link-icon {cor}".strip()
    return (
        f'      <a class="link" href="{html.escape(href)}" target="_blank" rel="noopener">\n'
        f'        <span class="{classe}">{ICONE}</span>\n'
        f'        <span class="link-text"><span class="link-title">{html.escape(titulo)}</span>'
        f'<span class="link-desc">{html.escape(descricao)}</span></span>\n'
        f'      </a>\n'
    )


def ler_config() -> dict:
    caminho = PASTA_HUB / CONFIG
    if caminho.exists():
        try:
            return json.loads(caminho.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            pass
    return {"site": "", "relatorios": {}}


def normalizar_site(texto: str) -> str:
    """Aceita 'hubadqrv', 'hubadqrv.netlify.app', 'https://hubadqrv.netlify.app/' ou o Site ID."""
    t = texto.strip().lower()
    t = re.sub(r"^https?://", "", t)
    t = t.split("/")[0]
    return t.replace(".netlify.app", "").strip()


def pedir_site(atual: str = "") -> str:
    print("Informe o endereço do seu site no Netlify (ex.: https://hubadqrv.netlify.app/)")
    print("ou o Project ID / Site ID que aparece em Site configuration > General no Netlify.")
    print("Dica: o Site ID não muda nem quando você troca o endereço do site.")
    site = normalizar_site(perguntar("Site", atual))
    if not site:
        sys.exit("Nenhum site informado.")
    return site


def gravar_config(cfg: dict) -> None:
    (PASTA_HUB / CONFIG).write_text(json.dumps(cfg, indent=2, ensure_ascii=False), encoding="utf-8")


def carregar_index():
    index = PASTA_HUB / "index.html"
    if not index.exists():
        sys.exit(f"Não encontrei o index.html na pasta do script: {PASTA_HUB}")
    pagina = index.read_text(encoding="utf-8")
    if not MARCA.search(pagina):
        sys.exit("O index.html não tem marcações de atalhos "
                 "(ATALHOS:INICIO / ATALHOS:FIM ou as dos submenus, como ATALHOS:COMERCIAL:INICIO).")
    return index, pagina


def menus_disponiveis() -> list:
    """Menus encontrados nas páginas da pasta, na ordem em que aparecem."""
    lista = []
    for nome_arq in PAGINAS:
        arq = PASTA_HUB / nome_arq
        if not arq.is_file():
            continue
        pagina = arq.read_text(encoding="utf-8")
        h1 = re.search(r"<h1[^>]*>(.*?)</h1>", pagina, re.S | re.I)
        nome_pagina = html.unescape(re.sub(r"<[^>]+>", "", h1.group(1))).strip() if h1 else arq.stem
        chaves = MARCA.findall(pagina)
        for chave in chaves:
            inicio = f"<!-- ATALHOS:{chave}:INICIO -->" if chave else INICIO
            fim = f"<!-- ATALHOS:{chave}:FIM -->" if chave else FIM
            if fim not in pagina:
                continue
            if chave:
                nome = f"{nome_pagina} · {ROTULOS.get(chave.upper(), chave.replace('_', ' ').title())}"
            else:
                nome = nome_pagina
            lista.append({"nome": nome, "arquivo": nome_arq, "inicio": inicio, "fim": fim})
    # Duas páginas com o mesmo título: mostra o arquivo para diferenciar
    nomes = [m["nome"] for m in lista]
    for m in lista:
        if nomes.count(m["nome"]) > 1:
            m["nome"] += f" ({m['arquivo']})"
    return lista


def escolher_menu(acao: str, excluir: dict = None) -> dict:
    menus = [m for m in menus_disponiveis()
             if not excluir or (m["arquivo"], m["inicio"]) != (excluir["arquivo"], excluir["inicio"])]
    if not menus:
        sys.exit("Nenhuma página com marcações de atalhos foi encontrada.")
    if len(menus) == 1:
        return menus[0]
    print(f"Em qual menu deseja {acao}?\n")
    for i, m in enumerate(menus, 1):
        print(f"  {i}. {m['nome']}")
    r = perguntar("\nEscolha (Enter cancela)")
    if not (r.isdigit() and 1 <= int(r) <= len(menus)):
        sys.exit("Nada foi alterado.")
    menu = menus[int(r) - 1]
    print(f"\nMenu: {menu['nome']}\n")
    return menu


def carregar_menu(menu: dict):
    arq = PASTA_HUB / menu["arquivo"]
    if not arq.is_file():
        sys.exit(f"Não encontrei o {menu['arquivo']} na pasta do script: {PASTA_HUB}")
    pagina = arq.read_text(encoding="utf-8")
    if menu["inicio"] not in pagina or menu["fim"] not in pagina:
        sys.exit(f"O {menu['arquivo']} não tem as marcações do menu {menu['nome']}.")
    return arq, pagina


def href_em_uso(href: str, ignorar_arquivo: Path = None, texto_atual: str = "") -> bool:
    """True se algum menu (de qualquer página) ainda aponta para esse relatório."""
    alvo = f'href="{html.escape(href)}"'
    if alvo in texto_atual:
        return True
    for nome in PAGINAS:
        arq = PASTA_HUB / nome
        if arq.is_file() and (ignorar_arquivo is None or arq.resolve() != ignorar_arquivo.resolve()):
            if alvo in arq.read_text(encoding="utf-8"):
                return True
    return False


def salvar(index: Path, pagina: str) -> None:
    backup = index.with_name(index.stem + "_backup.html")
    shutil.copy2(index, backup)
    index.write_text(pagina, encoding="utf-8")
    print(f"Cópia de segurança da versão anterior: {backup.name}")


# ------------------------------------------------------------------- incluir

def incluir() -> None:
    menu = escolher_menu("incluir os relatórios")
    index, pagina = carregar_menu(menu)
    cfg = ler_config()

    pasta_origem = limpar_caminho(perguntar("Endereço da pasta com os relatórios"))
    if not pasta_origem.is_dir():
        sys.exit(f"Pasta não encontrada: {pasta_origem}")

    relatorios = sorted(p for p in pasta_origem.iterdir()
                        if p.is_file() and p.suffix.lower() in (".html", ".htm")
                        and p.resolve() != index.resolve())
    if not relatorios:
        sys.exit("Nenhum arquivo .html encontrado nessa pasta.")

    destino = PASTA_HUB / PASTA_DESTINO
    destino.mkdir(exist_ok=True)

    qtd_atual = pagina.count('class="link"')
    novos = []
    print(f"\n{len(relatorios)} relatório(s) encontrado(s).\n")

    for arq in relatorios:
        print(f"→ {arq.name}")
        if not sim("  Incluir no menu?"):
            print("  Ignorado.\n")
            continue

        titulo = perguntar("  Nome no menu", titulo_do_html(arq))
        descricao = perguntar("  Descrição curta", "")
        nome_final = slug(titulo) + ".html"
        href = f"{PASTA_DESTINO}/{nome_final}"

        if f'href="{href}"' in pagina:
            if sim("  Já existe um atalho com esse nome. Substituir o arquivo?", "n"):
                shutil.copy2(arq, destino / nome_final)
                cfg["relatorios"][href] = str(arq.resolve())
                print("  Arquivo atualizado. O atalho continua o mesmo.\n")
            else:
                print("  Ignorado.\n")
            continue

        shutil.copy2(arq, destino / nome_final)
        cfg["relatorios"][href] = str(arq.resolve())
        cor = CORES[(qtd_atual + len(novos)) % len(CORES)]
        novos.append(montar_atalho(href, titulo, descricao, cor))
        print(f"  Incluído: {href}\n")

    gravar_config(cfg)
    if not novos:
        print("Nenhum atalho novo adicionado.")
        return

    # Insere antes da linha da marcação FIM, qualquer que seja o recuo dela
    pos = pagina.index(menu["fim"])
    inicio_linha = pagina.rfind("\n", 0, pos) + 1
    if pagina[inicio_linha:pos].strip():  # marcação no meio de uma linha
        inicio_linha = pos
    pagina = pagina[:inicio_linha] + "".join(novos) + pagina[inicio_linha:]
    salvar(index, pagina)

    # Confere se os atalhos foram mesmo gravados
    gravado = index.read_text(encoding="utf-8")
    faltando = [a for a in novos if re.search(r'href="([^"]+)"', a).group(1) not in gravado]
    if faltando:
        sys.exit(f"ERRO: os atalhos não foram gravados no {index.name}. Nada foi publicado.")
    print(f"{len(novos)} atalho(s) adicionado(s) ao menu {menu['nome']}.")
    print(f"Arquivo alterado: {index}")
    print("Abra esse arquivo no navegador para conferir (Ctrl+F5 se já estiver aberto).")
    print("Para o atalho aparecer no site, use a opção 3 para publicar no Netlify.")


# ------------------------------------------------------------ atalhos do menu

def atalhos_do_index(pagina: str, inicio: str = INICIO, fim: str = FIM) -> list:
    """Lê os atalhos que estão entre as marcações do menu."""
    ini = pagina.index(inicio) + len(inicio)
    bloco = pagina[ini:pagina.index(fim)]
    lista = []
    for m in re.finditer(r'[ \t]*<a class="link" href="([^"]*)".*?</a>[ \t]*\n?', bloco, re.S):
        t = re.search(r'class="link-title">(.*?)</span>', m.group(0), re.S)
        d = re.search(r'class="link-desc">(.*?)</span>', m.group(0), re.S)
        href = html.unescape(m.group(1))
        lista.append({
            "href": href,
            "titulo": html.unescape(t.group(1)).strip() if t else href,
            "descricao": html.unescape(d.group(1)).strip() if d else "",
            "externo": href.startswith(("http://", "https://")),
            "match": m,
        })
    return lista


# ------------------------------------------------------------------- excluir

def excluir() -> None:
    menu = escolher_menu("excluir atalhos")
    index, pagina = carregar_menu(menu)
    cfg = ler_config()

    ini = pagina.index(menu["inicio"]) + len(menu["inicio"])
    fim = pagina.index(menu["fim"])
    bloco = pagina[ini:fim]
    itens = atalhos_do_index(pagina, menu["inicio"], menu["fim"])
    atalhos = [a["match"] for a in itens]
    if not atalhos:
        print("O menu não tem atalhos para excluir.")
        return

    print("Atalhos no menu:\n")
    for i, a in enumerate(itens, 1):
        print(f"  {i}. {a['titulo']}  ({a['href']})")

    resposta = perguntar("\nNúmeros dos atalhos a excluir, separados por vírgula (Enter cancela)")
    escolhidos = {int(p) - 1 for p in re.split(r"[,\s]+", resposta)
                  if p.isdigit() and 1 <= int(p) <= len(atalhos)}
    if not escolhidos:
        print("Nada foi excluído.")
        return

    print("\nSerão excluídos:")
    for i in sorted(escolhidos):
        print(f"  - {html.unescape(atalhos[i].group(1))}")
    if not sim("Confirmar?", "n"):
        print("Nada foi excluído.")
        return

    novo_bloco = bloco
    for i in sorted(escolhidos, reverse=True):
        m = atalhos[i]
        novo_bloco = novo_bloco[:m.start()] + novo_bloco[m.end():]
    for i in sorted(escolhidos):
        href = html.unescape(atalhos[i].group(1))
        if href_em_uso(href, index, pagina[:ini] + novo_bloco + pagina[fim:]):
            print(f"  {href} continua em outro menu: o arquivo foi mantido.")
            continue
        cfg["relatorios"].pop(href, None)
        arquivo = PASTA_HUB / href
        if href.startswith(PASTA_DESTINO + "/") and arquivo.is_file():
            if sim(f"  Apagar também o arquivo {href}?"):
                arquivo.unlink()
                print(f"  Arquivo apagado: {href}")
    if not novo_bloco.startswith("\n"):
        novo_bloco = "\n" + novo_bloco
    if not novo_bloco.endswith("\n"):
        novo_bloco += "\n"

    gravar_config(cfg)
    pagina = pagina[:ini] + novo_bloco + pagina[fim:]
    print(f"\n{len(escolhidos)} atalho(s) excluído(s) do menu.")
    salvar(index, pagina)
    print("Use a opção 3 para publicar no Netlify.")


# ------------------------------------------------------------------ usuários

USU_INICIO = "/* USUARIOS:INICIO */"
USU_FIM = "/* USUARIOS:FIM */"
CABECALHOS = {"usuario", "usuário", "usuarios", "usuários", "login", "user", "nome"}


def digitais_da_planilha(caminho: Path) -> list:
    """Lê usuário (coluna A) e senha (coluna B) e devolve as impressões digitais SHA-256."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("Falta o openpyxl. Instale com:  pip install openpyxl")
    planilha = load_workbook(caminho, read_only=True, data_only=True).active
    digitais, usuarios = [], []
    for i, linha in enumerate(planilha.iter_rows(values_only=True)):
        if not linha or len(linha) < 2 or linha[0] is None or linha[1] is None:
            continue
        usuario = str(linha[0]).strip().lower()
        senha = linha[1]
        if isinstance(senha, float) and senha.is_integer():
            senha = int(senha)
        senha = str(senha)
        if i == 0 and usuario in CABECALHOS:
            continue  # linha de título
        if not usuario or not senha:
            continue
        usuarios.append(usuario)
        digitais.append(hashlib.sha256(f"hubadq|{usuario}|{senha}".encode("utf-8")).hexdigest())
    return digitais, usuarios


def aplicar_usuarios(caminho: Path) -> bool:
    """Grava os usuários da planilha em todas as páginas do site. Retorna True se algo mudou."""
    carregar_index()
    digitais, usuarios = digitais_da_planilha(caminho)
    if not digitais:
        sys.exit("Nenhum usuário encontrado na planilha (usuário na coluna A, senha na coluna B).")
    novo = "\n" + ",\n".join(f'"{d}"' for d in digitais) + "\n"
    mudou = False
    for nome in PAGINAS:
        arq = PASTA_HUB / nome
        if not arq.is_file():
            continue
        pagina = arq.read_text(encoding="utf-8")
        if USU_INICIO not in pagina or USU_FIM not in pagina:
            if nome == "index.html":
                sys.exit("O index.html não tem as marcações de usuários (USUARIOS:INICIO / USUARIOS:FIM).")
            registrar(f"AVISO: {nome} não tem as marcações de usuários; o login dessa página não foi atualizado.")
            continue
        ini = pagina.index(USU_INICIO) + len(USU_INICIO)
        fim = pagina.index(USU_FIM)
        if pagina[ini:fim].strip() == novo.strip():
            continue
        salvar(arq, pagina[:ini] + novo + pagina[fim:])
        mudou = True
    if mudou:
        registrar(f"Usuários do login atualizados: {len(usuarios)} ({', '.join(usuarios)}).")
    else:
        registrar(f"Usuários do login sem alteração ({len(usuarios)}).")
    return mudou


def atualizar_usuarios() -> None:
    cfg = ler_config()
    caminho = limpar_caminho(perguntar("Endereço da planilha de usuários (.xlsx)", cfg.get("planilha_usuarios", "")))
    if not caminho.is_file():
        sys.exit(f"Planilha não encontrada: {caminho}")
    cfg["planilha_usuarios"] = str(caminho.resolve())
    gravar_config(cfg)
    aplicar_usuarios(caminho)
    print("A planilha ficou registrada: a publicação (opção 3 e automática) vai recarregá-la sempre.")
    print("Use a opção 3 para publicar no Netlify.")


# ---------------------------------------------------------------- publicar

def atualizar_copias(cfg: dict) -> int:
    """Copia para o HUB as versões novas dos relatórios. Retorna quantos mudaram."""
    alterados = 0
    for href, origem in cfg.get("relatorios", {}).items():
        origem = Path(origem)
        destino = PASTA_HUB / href
        if not origem.is_file():
            registrar(f"AVISO: arquivo de origem não encontrado: {origem}")
            continue
        if destino.is_file() and filecmp.cmp(origem, destino, shallow=False):
            continue
        destino.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(origem, destino)
        alterados += 1
        registrar(f"Atualizado: {href}")
    return alterados


def assinatura_conteudo() -> str:
    """Impressão digital do que vai para o site, para saber se algo mudou."""
    h = hashlib.sha256()
    arquivos = [PASTA_HUB / n for n in PAGINAS if (PASTA_HUB / n).is_file()]
    rel = PASTA_HUB / PASTA_DESTINO
    if rel.is_dir():
        arquivos += sorted(p for p in rel.rglob("*") if p.is_file())
    for arq in arquivos:
        h.update(str(arq.relative_to(PASTA_HUB)).encode())
        h.update(arq.read_bytes())
    return h.hexdigest()


def montar_pasta_publicacao() -> Path:
    """Monta uma pasta só com o que vai para o site (sem script, backup e config)."""
    # Fica na pasta temporária do Windows, fora do OneDrive, para a sincronização não atrapalhar.
    # Cada envio usa uma pasta nova, então sobras de envios anteriores não atrapalham.
    temp = Path(tempfile.gettempdir())
    for antiga in temp.glob("hub_adq_publicar*"):
        shutil.rmtree(antiga, ignore_errors=True)  # limpeza do que for possível apagar
    pub = Path(tempfile.mkdtemp(prefix="hub_adq_publicar_"))
    for nome in PAGINAS:
        if (PASTA_HUB / nome).is_file():
            shutil.copy2(PASTA_HUB / nome, pub / nome)
    rel = PASTA_HUB / PASTA_DESTINO
    if rel.is_dir():
        shutil.copytree(rel, pub / PASTA_DESTINO, dirs_exist_ok=True)
    return pub


def enviar_netlify(pub: Path, site: str):
    """Retorna (sucesso, site_nao_encontrado)."""
    netlify = shutil.which("netlify")
    if not netlify:
        registrar("ERRO: Netlify CLI não encontrado. Instale com: npm install -g netlify-cli")
        return False, False

    mensagem = "HUB ADQ " + datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
    base = [netlify, "deploy", "--prod", "--dir", str(pub), "--site", site, "--message", mensagem]

    ambiente = dict(os.environ, NO_COLOR="1", FORCE_COLOR="0")  # log sem códigos de cor
    for extra in (["--no-build"], []):  # versões antigas do CLI não têm --no-build
        # cwd: o Netlify CLI cria uma pasta .netlify onde é executado; sem isso,
        # rodando pelo Agendador ou como administrador, ele tenta usar C:\Windows\System32
        try:
            r = subprocess.run(base + extra, capture_output=True, text=True,
                               encoding="utf-8", errors="replace", cwd=str(PASTA_HUB),
                               env=ambiente, stdin=subprocess.DEVNULL, timeout=900)
        except subprocess.TimeoutExpired:
            registrar("ERRO: o Netlify não respondeu em 15 minutos. Verifique a internet e o login (netlify login).")
            return False, False
        saida = (r.stdout or "") + (r.stderr or "")
        if r.returncode != 0 and "unknown option" in saida.lower() and extra:
            continue
        break

    if r.returncode == 0:
        link = re.search(r"https://\S+\.netlify\.app\S*", saida)
        registrar("Publicado com sucesso" + (f": {link.group(0)}" if link else "."))
        return True, False

    registrar("ERRO na publicação. Mensagem do Netlify:")
    registrar(saida.strip()[-1500:])
    baixo = saida.lower()
    nao_encontrado = "not found" in baixo or "404" in baixo or "no site" in baixo or "no project" in baixo
    if nao_encontrado:
        registrar(f"O Netlify não encontrou o site '{site}'. Se você mudou o endereço do site, "
                  "use a opção 9 para informar o novo.")
    elif "credit usage exceeded" in baixo or "credits are added" in baixo:
        registrar("A conta do Netlify logada neste computador está sem créditos. Se quiser publicar "
                  "por outra conta, use a opção 10 para entrar nela.")
    elif "login" in baixo or "auth" in baixo or "unauthorized" in baixo or "401" in baixo or "403" in baixo:
        registrar("Dica: use a opção 10 para entrar na conta do Netlify dona do site.")
    return False, nao_encontrado


def publicar() -> None:
    carregar_index()
    cfg = ler_config()

    if not cfg.get("site"):
        if AUTOMATICO:
            sys.exit("Site do Netlify não configurado. Rode o script manualmente e use a opção 3 uma vez.")
        cfg["site"] = pedir_site()
        gravar_config(cfg)

    registrar(f"Iniciando atualização do HUB ({cfg['site']})")
    planilha = cfg.get("planilha_usuarios")
    if planilha:
        if Path(planilha).is_file():
            aplicar_usuarios(Path(planilha))
        else:
            registrar(f"AVISO: planilha de usuários não encontrada ({planilha}). Mantidos os usuários atuais.")
    alterados = atualizar_copias(cfg)
    registrar(f"{alterados} relatório(s) com versão nova.")

    assinatura = assinatura_conteudo()
    if assinatura == cfg.get("ultima_publicacao"):
        if AUTOMATICO and "--sempre" not in sys.argv:
            registrar("O site já está com a versão mais recente. Publicação dispensada.")
            return
        if not AUTOMATICO and not sim("O site já está com a versão mais recente. Publicar mesmo assim?", "n"):
            return

    antiga = PASTA_HUB / PASTA_PUBLICACAO  # pasta usada por versões anteriores do script
    if antiga.exists():
        shutil.rmtree(antiga, ignore_errors=True)
    registrar("Montando os arquivos para envio...")
    pub = montar_pasta_publicacao()
    registrar("Enviando para o Netlify (pode levar alguns minutos)...")
    sucesso, nao_encontrado = enviar_netlify(pub, cfg["site"])
    if not sucesso and nao_encontrado and not AUTOMATICO:
        print()
        if sim("Quer informar o endereço novo do site e tentar de novo?", "s"):
            cfg["site"] = pedir_site()
            gravar_config(cfg)
            registrar(f"Site do Netlify alterado para: {cfg['site']}")
            registrar("Enviando para o Netlify (pode levar alguns minutos)...")
            sucesso, _ = enviar_netlify(pub, cfg["site"])
    shutil.rmtree(pub, ignore_errors=True)
    if sucesso:
        cfg["ultima_publicacao"] = assinatura
        gravar_config(cfg)
    elif AUTOMATICO:
        sys.exit(1)


def trocar_site() -> None:
    cfg = ler_config()
    atual = cfg.get("site", "")
    print(f"Site atual: {atual or '(nenhum)'}\n")
    novo = pedir_site(atual)
    if novo == atual:
        print("Nada foi alterado.")
        return
    cfg["site"] = novo
    cfg.pop("ultima_publicacao", None)  # força a próxima publicação a enviar tudo
    gravar_config(cfg)
    registrar(f"Site do Netlify alterado: {atual or '(nenhum)'} -> {novo}")
    print("Use a opção 3 para publicar no site novo.")


def caminho_chrome() -> str:
    """Procura o Google Chrome instalado no Windows (ou no PATH)."""
    candidatos = []
    for var in ("PROGRAMFILES", "PROGRAMFILES(X86)", "LOCALAPPDATA"):
        base = os.environ.get(var)
        if base:
            candidatos.append(Path(base) / "Google" / "Chrome" / "Application" / "chrome.exe")
    for c in candidatos:
        if c.is_file():
            return str(c)
    for nome in ("chrome", "google-chrome", "chrome.exe"):
        achado = shutil.which(nome)
        if achado:
            return achado
    return ""


def abrir_no_chrome(url: str) -> bool:
    chrome = caminho_chrome()
    cfg = ler_config()
    perfil = cfg.get("chrome_perfil")  # opcional, ex.: "Profile 1" (veja chrome://version)
    if chrome:
        cmd = [chrome] + ([f"--profile-directory={perfil}"] if perfil else []) + [url]
        try:
            subprocess.Popen(cmd)
            return True
        except OSError:
            pass
    if os.name == "nt":  # última tentativa: deixa o Windows achar o Chrome
        try:
            subprocess.Popen(["cmd", "/c", "start", "", "chrome", url])
            return True
        except OSError:
            pass
    return False


def rodar_netlify(args: list, timeout: int = 120) -> str:
    """Roda o Netlify CLI e devolve a saída (texto)."""
    ambiente = dict(os.environ, NO_COLOR="1", FORCE_COLOR="0", BROWSER="none")
    try:
        r = subprocess.run([shutil.which("netlify")] + args, capture_output=True, text=True,
                           encoding="utf-8", errors="replace", cwd=str(PASTA_HUB), env=ambiente,
                           stdin=subprocess.DEVNULL, timeout=timeout)
    except subprocess.TimeoutExpired:
        return ""
    return (r.stdout or "") + (r.stderr or "")


def json_da_saida(saida: str) -> dict:
    ini, fim = saida.find("{"), saida.rfind("}")
    if ini >= 0 and fim > ini:
        try:
            return json.loads(saida[ini:fim + 1])
        except json.JSONDecodeError:
            pass
    return {}


def entrar_netlify() -> None:
    if not shutil.which("netlify"):
        sys.exit("Netlify CLI não encontrado. Instale com: npm install -g netlify-cli")
    if os.environ.get("NETLIFY_AUTH_TOKEN"):
        print("ATENÇÃO: a variável NETLIFY_AUTH_TOKEN está definida neste computador e tem")
        print("prioridade sobre o login. Apague-a nas Variáveis de Ambiente do Windows.\n")

    print("Vou abrir a autorização do Netlify no Google Chrome. Lá, confira se está na")
    print("conta certa (canto superior direito) e clique em Authorize.\n")
    if not sim("Continuar?", "s"):
        return

    # 1) Pede um "ticket" de login ao Netlify (não abre navegador nenhum)
    saida = rodar_netlify(["login", "--request", "HUB ADQ - incluir_relatorios.py", "--json"])
    if "unknown option" in saida.lower():
        sys.exit("Seu Netlify CLI é antigo. Atualize com:  npm install -g netlify-cli  e tente de novo.")
    dados = json_da_saida(saida)
    ticket = dados.get("ticket_id") or (re.search(r"Ticket ID:\s*(\S+)", saida) or [None, None])[1]
    url = dados.get("url") or (re.search(r"https://app\.netlify\.com/authorize\S+", saida) or [None])[0]
    if not ticket or not url:
        registrar("ERRO ao iniciar o login. Mensagem do Netlify:")
        registrar(saida.strip()[-1500:] or "(sem resposta; verifique a internet)")
        return

    # 2) Abre o link no Chrome
    if abrir_no_chrome(url):
        print("\nAbri o Chrome. Autorize por lá; o script percebe sozinho.")
    else:
        print("\nNão achei o Chrome. Copie o endereço abaixo e cole no Chrome:\n")
        print("  " + url)
    print("Aguardando a autorização (até 10 minutos; Ctrl+C cancela)...")

    # 3) Confere a cada 3 segundos se a autorização foi feita
    limite = time.time() + 600
    while time.time() < limite:
        time.sleep(3)
        saida = rodar_netlify(["login", "--check", ticket, "--json"], timeout=60)
        dados = json_da_saida(saida)
        status = dados.get("status") or (re.search(r"Status:\s*(\w+)", saida) or [None, ""])[1]
        if status == "authorized":
            user = dados.get("user") or {}
            nome = user.get("name") or ""
            email = user.get("email") or (re.search(r"Email:\s*(\S+)", saida) or [None, ""])[1]
            print(f"\nPronto! Conectado como: {nome} {('<' + email + '>') if email else ''}".rstrip())
            registrar(f"Login no Netlify refeito pela opção 10: {email or nome}")
            print("Se o site ainda for o da conta antiga, use a opção 9 para informar o site da conta nova.")
            return
        if status == "denied":
            sys.exit("A autorização foi recusada ou expirou. Use a opção 10 de novo.")
    sys.exit("A autorização não foi feita em 10 minutos. Use a opção 10 de novo.")


# ------------------------------------------------------------------- ordenar

def chave_alfabetica(texto: str) -> str:
    return unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode().lower()


def ordenar() -> None:
    menu = escolher_menu("ordenar os relatórios")
    index, pagina = carregar_menu(menu)
    itens = atalhos_do_index(pagina, menu["inicio"], menu["fim"])
    if len(itens) < 2:
        print("O menu precisa ter pelo menos dois atalhos para ordenar.")
        return

    print("Ordem atual:\n")
    for i, a in enumerate(itens, 1):
        print(f"  {i}. {a['titulo']}")
    print("\nComo deseja ordenar?\n")
    print("  1. Ordem alfabética (A a Z)")
    print("  2. Mover um relatório para outra posição")
    print("  3. Digitar a nova ordem completa")
    modo = perguntar("\nEscolha (Enter cancela)")
    n = len(itens)
    ordem = list(range(n))

    if modo == "1":
        ordem.sort(key=lambda i: chave_alfabetica(itens[i]["titulo"]))
    elif modo == "2":
        de = perguntar("Número do relatório que vai mudar de lugar")
        para = perguntar(f"Nova posição (1 a {n})")
        if not (de.isdigit() and para.isdigit() and 1 <= int(de) <= n and 1 <= int(para) <= n):
            sys.exit("Números inválidos. Nada foi alterado.")
        ordem.remove(int(de) - 1)
        ordem.insert(int(para) - 1, int(de) - 1)
    elif modo == "3":
        print("Digite os números na ordem desejada, separados por vírgula (ex.: 3,1,2).")
        print("Os que você não citar ficam no final, na ordem em que estão hoje.")
        numeros = [int(x) - 1 for x in re.split(r"[,\s]+", perguntar("Nova ordem")) if x.isdigit()]
        invalidos = [x + 1 for x in numeros if not 0 <= x < n]
        if invalidos or len(set(numeros)) != len(numeros) or not numeros:
            sys.exit("Ordem inválida (número repetido, fora da lista ou vazio). Nada foi alterado.")
        ordem = numeros + [i for i in range(n) if i not in numeros]
    else:
        print("Nada foi alterado.")
        return

    if ordem == list(range(n)):
        print("\nA ordem já é essa. Nada foi alterado.")
        return

    print("\nNova ordem:\n")
    for pos, i in enumerate(ordem, 1):
        print(f"  {pos}. {itens[i]['titulo']}")
    if not sim("\nConfirmar?"):
        print("Nada foi alterado.")
        return

    ini = pagina.index(menu["inicio"]) + len(menu["inicio"])
    fim = pagina.index(menu["fim"])
    bloco = pagina[ini:fim]
    primeiro, ultimo = itens[0]["match"], itens[-1]["match"]
    trechos = [itens[i]["match"].group(0).rstrip() + "\n" for i in ordem]
    novo_bloco = bloco[:primeiro.start()] + "".join(trechos) + bloco[ultimo.end():].lstrip("\n")
    salvar(index, pagina[:ini] + novo_bloco + pagina[fim:])

    # Confere a gravação
    gravados = [a["href"] for a in atalhos_do_index(index.read_text(encoding="utf-8"), menu["inicio"], menu["fim"])]
    if gravados != [itens[i]["href"] for i in ordem]:
        sys.exit(f"ERRO: a nova ordem não foi gravada corretamente. Restaure o {index.stem}_backup.html.")
    print("Ordem atualizada. Use a opção 3 para publicar no Netlify.")


# --------------------------------------------------------------------- mover

def mover() -> None:
    origem = escolher_menu("mover relatórios (menu de origem)")
    arq_o, pag_o = carregar_menu(origem)
    itens = atalhos_do_index(pag_o, origem["inicio"], origem["fim"])
    if not itens:
        print("Esse menu não tem atalhos para mover.")
        return

    print("Atalhos no menu:\n")
    for i, a in enumerate(itens, 1):
        print(f"  {i}. {a['titulo']}")
    resposta = perguntar("\nNúmeros dos atalhos a mover, separados por vírgula (Enter cancela)")
    escolhidos = sorted({int(p) - 1 for p in re.split(r"[,\s]+", resposta)
                         if p.isdigit() and 1 <= int(p) <= len(itens)})
    if not escolhidos:
        print("Nada foi alterado.")
        return

    print()
    destino = escolher_menu("colocar esses relatórios (menu de destino)", excluir=origem)
    manter = False
    if destino["arquivo"] != origem["arquivo"]:
        manter = sim(f"Manter também no menu {origem['nome']}? (n = mover, s = copiar)", "n")

    acao = "copiados" if manter else "movidos"
    print(f"\nSerão {acao} de {origem['nome']} para {destino['nome']}:")
    for i in escolhidos:
        print(f"  - {itens[i]['titulo']}")
    if not sim("Confirmar?", "n"):
        print("Nada foi alterado.")
        return

    # Trabalha com o texto das páginas na memória e grava cada arquivo uma vez só
    textos = {arq_o: pag_o}

    # 1) Retira do menu de origem
    if not manter:
        ini = pag_o.index(origem["inicio"]) + len(origem["inicio"])
        fim = pag_o.index(origem["fim"])
        bloco = pag_o[ini:fim]
        for i in reversed(escolhidos):
            m = itens[i]["match"]
            bloco = bloco[:m.start()] + bloco[m.end():]
        if not bloco.startswith("\n"):
            bloco = "\n" + bloco
        if not bloco.endswith("\n"):
            bloco += "\n"
        textos[arq_o] = pag_o[:ini] + bloco + pag_o[fim:]

    # 2) Coloca no menu de destino (antes da marcação FIM)
    arq_d = PASTA_HUB / destino["arquivo"]
    if arq_d not in textos:
        textos[arq_d] = carregar_menu(destino)[1]
    pag_d = textos[arq_d]
    ja_no_destino = {a["href"] for a in atalhos_do_index(pag_d, destino["inicio"], destino["fim"])}
    trechos = []
    for i in escolhidos:
        if itens[i]["href"] in ja_no_destino:
            print(f"  {itens[i]['titulo']} já estava em {destino['nome']}.")
            continue
        trechos.append(itens[i]["match"].group(0).rstrip() + "\n")
    if trechos:
        pos = pag_d.index(destino["fim"])
        inicio_linha = pag_d.rfind("\n", 0, pos) + 1
        if pag_d[inicio_linha:pos].strip():
            inicio_linha = pos
        textos[arq_d] = pag_d[:inicio_linha] + "".join(trechos) + pag_d[inicio_linha:]

    originais = {arq_o: pag_o, arq_d: pag_d}
    for arq, texto in textos.items():
        if texto != originais.get(arq):
            salvar(arq, texto)

    # Confere a gravação
    gravados = {a["href"] for a in atalhos_do_index(arq_d.read_text(encoding="utf-8"),
                                                     destino["inicio"], destino["fim"])}
    faltando = [itens[i]["titulo"] for i in escolhidos if itens[i]["href"] not in gravados]
    if faltando:
        sys.exit("ERRO: não foram gravados no destino: " + ", ".join(faltando)
                 + f". Restaure o {arq_d.stem}_backup.html.")
    print(f"\n{len(escolhidos)} relatório(s) {acao} para {destino['nome']}.")
    print("Use a opção 3 para publicar no Netlify.")


# ------------------------------------------------------------- ver e validar

def ver_relatorios() -> None:
    carregar_index()
    for menu in menus_disponiveis():
        _, pagina = carregar_menu(menu)
        print(f"=== {menu['nome']} ({menu['arquivo']}) ===\n")
        ver_menu(atalhos_do_index(pagina, menu["inicio"], menu["fim"]))


def ver_menu(itens: list) -> None:
    cfg = ler_config()
    if not itens:
        print("  Nenhum atalho neste menu.\n")
        return
    site = cfg.get("site")
    print(f"{len(itens)} relatório(s) no menu:\n")
    for i, a in enumerate(itens, 1):
        print(f"  {i}. {a['titulo']}")
        if a["descricao"]:
            print(f"     {a['descricao']}")
        if a["externo"]:
            print(f"     Link externo: {a['href']}")
        else:
            print(f"     Arquivo no HUB: {a['href']}")
            if site:
                print(f"     No site:        https://{site}.netlify.app/{a['href']}")
            origem = cfg.get("relatorios", {}).get(a["href"])
            print(f"     Origem:         {origem or '(não registrada — inclua de novo pela opção 1)'}")
        print()


def baixar(url: str, timeout: int = 20):
    """Retorna (código HTTP, conteúdo) ou (None, mensagem de erro)."""
    pedido = urllib.request.Request(url, headers={"User-Agent": "HUB-ADQ-validador/1.0"})
    try:
        with urllib.request.urlopen(pedido, timeout=timeout) as r:
            return r.status, r.read()
    except urllib.error.HTTPError as e:
        return e.code, b""
    except Exception as e:  # sem internet, proxy, certificado etc.
        return None, str(getattr(e, "reason", e))


def recursos_locais_faltando(arquivo: Path) -> list:
    """Arquivos que o relatório chama (imagens, scripts, estilos) e que não estão ao lado dele."""
    texto = arquivo.read_text(encoding="utf-8", errors="ignore")
    faltando = []
    for ref in re.findall(r'(?:src|href)\s*=\s*["\']([^"\'#?]+)', texto, re.I):
        ref = ref.strip()
        if not ref or re.match(r"^(https?:|data:|mailto:|tel:|javascript:|//|blob:)", ref, re.I):
            continue
        if re.match(r"^[a-zA-Z]:[\\/]|^file:", ref):
            faltando.append(f"{ref} (caminho do seu computador; não funciona no site)")
        elif not (arquivo.parent / ref).exists():
            faltando.append(ref)
    return sorted(set(faltando))


def validar() -> None:
    carregar_index()
    cfg = ler_config()
    itens = []
    for menu in menus_disponiveis():
        _, pagina = carregar_menu(menu)
        for a in atalhos_do_index(pagina, menu["inicio"], menu["fim"]):
            a["menu"] = menu["nome"]
            itens.append(a)
    site = cfg.get("site")
    if not itens:
        print("O menu não tem atalhos.")
        return

    online = sim("Verificar também o site no ar (precisa de internet)?")
    print()
    total_erros = total_avisos = 0

    menu_atual = None
    for i, a in enumerate(itens, 1):
        if a["menu"] != menu_atual:
            menu_atual = a["menu"]
            print(f"=== {menu_atual} ===\n")
        erros, avisos, oks = [], [], []
        if a["externo"]:
            if online:
                cod, _ = baixar(a["href"])
                if cod == 200:
                    oks.append("link externo responde normalmente")
                elif cod is None:
                    avisos.append(f"não foi possível acessar o link ({_})")
                else:
                    erros.append(f"link externo respondeu com erro {cod}")
            else:
                oks.append("link externo (não verificado sem internet)")
        else:
            arq = PASTA_HUB / a["href"]
            if not arq.is_file():
                erros.append(f"arquivo não existe no HUB: {a['href']}")
            else:
                conteudo = arq.read_bytes()
                inicio = conteudo[:4000].decode("utf-8", errors="ignore").lower()
                if not conteudo.strip():
                    erros.append("arquivo está vazio")
                elif "<html" not in inicio and "<body" not in inicio and "<!doctype" not in inicio:
                    avisos.append("o arquivo não parece ser uma página HTML")
                else:
                    tam = len(conteudo)
                    oks.append(f"arquivo no HUB ok ({tam / 1024:.0f} KB)" if tam >= 1024 else f"arquivo no HUB ok ({tam} bytes)")

                origem = cfg.get("relatorios", {}).get(a["href"])
                if not origem:
                    avisos.append("origem não registrada: a atualização diária não vai alcançar este relatório (inclua de novo pela opção 1)")
                elif not Path(origem).is_file():
                    erros.append(f"arquivo de origem não encontrado: {origem}")
                elif not filecmp.cmp(origem, arq, shallow=False):
                    avisos.append("a origem tem versão mais nova que a cópia do HUB (rode a opção 3)")
                else:
                    oks.append("cópia do HUB igual à origem")

                for falta in recursos_locais_faltando(arq):
                    avisos.append(f"o relatório chama um arquivo que não vai para o site: {falta}")

                if online and site:
                    url = f"https://{site}.netlify.app/{a['href']}"
                    cod, corpo = baixar(url)
                    if cod == 200:
                        if hashlib.sha256(corpo).digest() == hashlib.sha256(conteudo).digest():
                            oks.append("no ar e atualizado")
                        else:
                            avisos.append("no ar, mas diferente da versão do HUB (rode a opção 3)")
                    elif cod is None:
                        avisos.append(f"não foi possível acessar o site ({corpo})")
                    else:
                        erros.append(f"no site, o endereço responde com erro {cod} (publique pela opção 3)")
                elif online and not site:
                    avisos.append("site do Netlify ainda não configurado (rode a opção 3 uma vez)")

        situacao = "ERRO " if erros else ("AVISO" if avisos else "OK   ")
        print(f"[{situacao}] {i}. {a['titulo']}  ({a['href']})")
        for m in erros:
            print(f"          ✗ {m}")
        for m in avisos:
            print(f"          ! {m}")
        for m in oks:
            print(f"          ✓ {m}")
        print()
        total_erros += bool(erros)
        total_avisos += bool(avisos) and not erros

    # Arquivos esquecidos na pasta relatorios, sem atalho no menu
    hrefs = {a["href"] for a in itens}
    pasta = PASTA_HUB / PASTA_DESTINO
    soltos = [f"{PASTA_DESTINO}/{p.name}" for p in sorted(pasta.glob("*.htm*"))] if pasta.is_dir() else []
    soltos = [h for h in soltos if h not in hrefs]
    if soltos:
        print("Arquivos na pasta relatorios sem atalho no menu (vão para o site, mas ninguém chega a eles):")
        for h in soltos:
            print(f"  - {h}")
        print()

    print(f"Resumo: {len(itens) - total_erros - total_avisos} ok, {total_avisos} com aviso, {total_erros} com erro.")


# ---------------------------------------------------------------------- main

def main() -> None:
    os.chdir(PASTA_HUB)
    if AUTOMATICO:
        publicar()
        return

    acoes = {
        "1": ("Incluir relatórios no menu", incluir),
        "2": ("Excluir atalhos do menu", excluir),
        "3": ("Atualizar relatórios e publicar no Netlify", publicar),
        "4": ("Atualizar usuários do login (planilha)", atualizar_usuarios),
        "5": ("Ver os relatórios do menu", ver_relatorios),
        "6": ("Validar se os relatórios estão funcionando", validar),
        "7": ("Ordenar os relatórios do menu", ordenar),
        "8": ("Mover relatórios para outro menu", mover),
        "9": ("Trocar o site do Netlify", trocar_site),
        "10": ("Entrar em outra conta do Netlify (Chrome)", entrar_netlify),
    }
    carregar_index()
    while True:
        print("\n=== HUB ADQ · Atalhos e publicação ===\n")
        for chave, (nome, _) in acoes.items():
            print(f"  {chave:>2}. {nome}")
        print("   0. Sair\n")
        try:
            opcao = perguntar("Escolha uma opção")
        except EOFError:
            return
        if opcao in ("0", "s", "sair"):
            return
        if opcao not in acoes:
            print("Opção inválida.")
            continue
        print(f"\n--- {acoes[opcao][0]} ---\n")
        try:
            acoes[opcao][1]()
        except KeyboardInterrupt:
            print("\nCancelado. Voltando ao menu.")
        except EOFError:
            return
        except SystemExit as erro:  # mensagens de erro das opções: mostra e volta ao menu
            if erro.code and erro.code != 1:
                print(f"\n{erro.code}")
        except Exception:
            registrar("ERRO inesperado. Detalhes técnicos:")
            registrar(traceback.format_exc())
        try:
            input("\nPressione Enter para voltar ao menu...")
        except EOFError:
            return


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nAté logo.")
    except SystemExit as erro:
        if erro.code and erro.code != 1:
            registrar(str(erro.code))
        if AUTOMATICO:
            raise
        input("\nPressione Enter para sair...")
    except Exception:
        registrar("ERRO inesperado. Detalhes técnicos:")
        registrar(traceback.format_exc())
        if AUTOMATICO:
            sys.exit(1)
        input("\nPressione Enter para sair...")
